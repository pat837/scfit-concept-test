import hmac
import os
import uuid
from collections import Counter
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from io import BytesIO

from dotenv import load_dotenv
from fastapi import Depends, FastAPI, Header, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import inspect, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

load_dotenv()

from .database import Base, SessionLocal, engine
from .models import SurveyResponse, SurveySession, utc_now
from .schemas import INITIAL_TESTING_INTEREST, USC_AFFILIATIONS, EventRequest, SessionRequest, SurveySubmission


LEGACY_REQUIRED_COLUMNS = (
    "fitness_frequency",
    "opportunity_difficulty",
    "partner_difficulty",
    "discovery_methods",
    "prototype_sections_explored",
    "community_value",
    "future_test_interest",
    "preferred_future_experience",
)


def migrate_response_schema() -> None:
    columns = {column["name"]: column for column in inspect(engine).get_columns("survey_responses")}
    with engine.begin() as connection:
        if "initial_product_testing_interest" not in columns:
            connection.execute(text(
                "ALTER TABLE survey_responses "
                "ADD COLUMN initial_product_testing_interest VARCHAR(40)"
            ))
        for column in LEGACY_REQUIRED_COLUMNS:
            if column in columns and not columns[column]["nullable"]:
                connection.execute(text(
                    f'ALTER TABLE survey_responses ALTER COLUMN "{column}" DROP NOT NULL'
                ))


@asynccontextmanager
async def lifespan(_: FastAPI):
    Base.metadata.create_all(bind=engine)
    migrate_response_schema()
    yield


app = FastAPI(title="SCFit Concept Test API", version="1.0.0", lifespan=lifespan)
origins = [origin.strip() for origin in os.getenv("FRONTEND_ORIGIN", "http://localhost:3000").split(",") if origin.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type", "X-Admin-Key"],
)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def require_admin(x_admin_key: str | None = Header(default=None)) -> None:
    expected = os.getenv("ADMIN_API_KEY", "")
    if not expected or not x_admin_key or not hmac.compare_digest(x_admin_key, expected):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid admin key.")


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/api/sessions")
def create_session(payload: SessionRequest, db: Session = Depends(get_db)):
    if payload.session_id:
        existing = db.get(SurveySession, payload.session_id)
        if existing:
            return {"session_id": existing.id, "created_at": existing.created_at}
    session = SurveySession(id=str(uuid.uuid4()))
    db.add(session)
    db.commit()
    return {"session_id": session.id, "created_at": session.created_at}


@app.post("/api/events/prototype-opened")
def prototype_opened(payload: EventRequest, db: Session = Depends(get_db)):
    session = _get_session(db, payload.session_id)
    if session.prototype_opened_at is None:
        session.prototype_opened_at = utc_now()
        db.commit()
    return {"status": "ok", "prototype_opened_at": session.prototype_opened_at}


@app.post("/api/events/prototype-returned")
def prototype_returned(payload: EventRequest, db: Session = Depends(get_db)):
    session = _get_session(db, payload.session_id)
    if session.prototype_opened_at is None:
        raise HTTPException(status_code=400, detail="Open the prototype before recording a return.")
    if session.prototype_returned_at is None:
        session.prototype_returned_at = utc_now()
        db.commit()
    return {"status": "ok", "prototype_returned_at": session.prototype_returned_at}


def _get_session(db: Session, session_id: str) -> SurveySession:
    session = db.get(SurveySession, session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Survey session not found. Start the concept test again.")
    return session


@app.post("/api/responses", status_code=status.HTTP_201_CREATED)
def create_response(payload: SurveySubmission, db: Session = Depends(get_db)):
    survey_session = _get_session(db, payload.session_id)
    if not survey_session.prototype_opened_at or not survey_session.prototype_returned_at:
        raise HTTPException(status_code=400, detail="Open the prototype and confirm your return before submitting.")
    if db.scalar(select(SurveyResponse.id).where(SurveyResponse.session_id == payload.session_id)):
        raise HTTPException(status_code=409, detail="This response has already been submitted.")

    now = utc_now()
    data = payload.model_dump(exclude={"session_id"})
    if data.get("email"):
        data["email"] = str(data["email"])
    response = SurveyResponse(
        session_id=payload.session_id,
        created_at=survey_session.created_at,
        submitted_at=now,
        prototype_opened_at=survey_session.prototype_opened_at,
        prototype_returned_at=survey_session.prototype_returned_at,
        completion_time_seconds=max(0, (now - _aware(survey_session.created_at)).total_seconds()),
        **data,
    )
    db.add(response)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="This response has already been submitted.")
    return {"status": "ok", "message": "Your response has been recorded."}


def _aware(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value


def _responses(db: Session) -> list[SurveyResponse]:
    return list(db.scalars(select(SurveyResponse).order_by(SurveyResponse.submitted_at)))


def _distribution(responses: list[SurveyResponse], field: str) -> dict[str, int]:
    counts = Counter()
    for response in responses:
        value = getattr(response, field)
        counts.update(value if isinstance(value, list) else [value] if value else [])
    if field == "initial_product_testing_interest":
        return {
            "Yes": counts[INITIAL_TESTING_INTEREST[0]],
            "Maybe": counts[INITIAL_TESTING_INTEREST[1]],
            "Not right now": counts[INITIAL_TESTING_INTEREST[2]],
        }
    return dict(counts.most_common())


@app.get("/api/admin/metrics", dependencies=[Depends(require_admin)])
def metrics(db: Session = Depends(get_db)):
    responses = _responses(db)
    total = len(responses)
    opened_count = sum(row.prototype_opened_at is not None for row in responses)
    return {
        "total_responses": total,
        "usc_undergraduate_count": sum(row.usc_affiliation == USC_AFFILIATIONS[0] for row in responses),
        "usc_graduate_count": sum(row.usc_affiliation == USC_AFFILIATIONS[1] for row in responses),
        "non_usc_count": sum(row.usc_affiliation == USC_AFFILIATIONS[3] for row in responses),
        "prototype_exploration_count": opened_count,
        "average_completion_time_seconds": round(sum(row.completion_time_seconds for row in responses) / total, 1) if total else 0,
        "distributions": {
            "initial_product_testing_interest": _distribution(responses, "initial_product_testing_interest"),
            "usc_affiliation": _distribution(responses, "usc_affiliation"),
            "fitness_interests": _distribution(responses, "fitness_interests"),
            "participation_barriers": _distribution(responses, "participation_barriers"),
            "most_valuable_feature": _distribution(responses, "most_valuable_feature"),
            "likely_action": _distribution(responses, "likely_action"),
            "adoption_barriers": _distribution(responses, "adoption_barriers"),
            "next_month_likelihood": _distribution(responses, "next_month_likelihood"),
            "overall_reaction": _distribution(responses, "overall_reaction"),
        },
    }


@app.get("/api/admin/export.xlsx", dependencies=[Depends(require_admin)])
def export_responses(db: Session = Depends(get_db)):
    responses = _responses(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Responses"
    fields = [
        ("Session ID", "session_id"), ("Participant Name", "participant_name"),
        ("Major / Program", "major_program"), ("College / School", "college_school"),
        ("Initial Product Testing Interest", "initial_product_testing_interest"),
        ("USC Affiliation", "usc_affiliation"), ("Fitness Interests", "fitness_interests"),
        ("Fitness Interests - Other", "fitness_interests_other"),
        ("Fitness Challenges", "participation_barriers"), ("Fitness Challenges - Other", "participation_barriers_other"),
        ("Feature That Stood Out", "most_valuable_feature"),
        ("Likely First Action", "likely_action"), ("Adoption Hesitations", "adoption_barriers"),
        ("Next-Month Likelihood", "next_month_likelihood"), ("Suggested Improvement", "overall_reaction"),
        ("Started At", "created_at"), ("Submitted At", "submitted_at"),
        ("Prototype Opened At", "prototype_opened_at"), ("Prototype Returned At", "prototype_returned_at"),
        ("Completion Time (seconds)", "completion_time_seconds"),
    ]
    sheet.append([label for label, _ in fields])
    for response in responses:
        values = []
        for _, field in fields:
            value = getattr(response, field)
            if isinstance(value, list):
                value = "; ".join(value)
            if isinstance(value, datetime):
                value = _aware(value).isoformat()
            values.append(_excel_safe(value))
        sheet.append(values)
    _format_sheet(sheet)

    summary = workbook.create_sheet("Summary")
    summary.append(["SCFit Concept Test Responses", "Value", "Response", "Count"])
    summary_rows = [
        ("Total participants", len(responses)),
        ("Prototype exploration count", sum(row.prototype_opened_at is not None for row in responses)),
        ("Average completion time (seconds)", round(sum(row.completion_time_seconds for row in responses) / len(responses), 1) if responses else 0),
    ]
    for label, value in summary_rows:
        summary.append([label, value])
    for title, field in [
        ("Initial product testing interest", "initial_product_testing_interest"),
        ("USC affiliation", "usc_affiliation"), ("Fitness interests", "fitness_interests"),
        ("Fitness challenges", "participation_barriers"), ("Feature that stood out", "most_valuable_feature"),
        ("Likely first action", "likely_action"), ("Adoption hesitations", "adoption_barriers"),
        ("Next-month usage likelihood", "next_month_likelihood"),
    ]:
        summary.append([title])
        for answer, count in _distribution(responses, field).items():
            summary.append([None, None, _excel_safe(answer), count])
    _format_sheet(summary)
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 20
    summary.column_dimensions["C"].width = 62
    summary.column_dimensions["D"].width = 12

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="scfit-concept-test-responses.xlsx"'},
    )


def _excel_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="204D3A")
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 48)
        sheet.column_dimensions[letter].width = width